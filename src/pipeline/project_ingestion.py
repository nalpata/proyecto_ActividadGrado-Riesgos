"""Ingesta en memoria de documentos nuevos por proyecto — Día 19A.

Reutiliza la configuración de recuperación seleccionada en los días 3 y 4:
chunking recursivo de 2.200 caracteres con solapamiento de 300 y embeddings
normalizados BAAI/bge-m3. No escribe documentos, texto ni vectores en disco.
"""

from __future__ import annotations

from hashlib import sha256
from io import BytesIO
from pathlib import Path
import re
from typing import Any
import unicodedata

import fitz
import numpy as np
from docx import Document

from src.preprocessing.clean_text import clean_text


EMBEDDING_MODEL = "BAAI/bge-m3"
CHUNK_SIZE_CHARS = 2200
CHUNK_OVERLAP_CHARS = 300
SUPPORTED_EXTENSIONS = {".pdf", ".docx"}
MAX_DOCUMENTS = 10
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TOTAL_BYTES = 50 * 1024 * 1024


def recursive_character_chunks(
    text: str,
    chunk_size: int = CHUNK_SIZE_CHARS,
    chunk_overlap: int = CHUNK_OVERLAP_CHARS,
) -> list[str]:
    """Divide texto con límites naturales y solapamiento reproducible."""

    value = clean_text(text)
    if not value:
        return []
    if chunk_size <= 0 or chunk_overlap < 0 or chunk_overlap >= chunk_size:
        raise ValueError("La configuración de chunking es inválida")

    separators = ("\n\n", "\n", ". ", "; ", " ")
    chunks: list[str] = []
    cursor = 0
    while cursor < len(value):
        hard_end = min(cursor + chunk_size, len(value))
        end = hard_end
        if hard_end < len(value):
            window = value[cursor:hard_end]
            minimum_boundary = max(1, chunk_size // 2)
            candidates = [window.rfind(separator) for separator in separators]
            boundary = max(candidates)
            if boundary >= minimum_boundary:
                end = cursor + boundary + 1
        chunk = value[cursor:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(value):
            break
        next_cursor = max(end - chunk_overlap, cursor + 1)
        while next_cursor < end and value[next_cursor].isspace():
            next_cursor += 1
        cursor = next_cursor
    return chunks


def _extract_pdf(content: bytes) -> list[tuple[int, str]]:
    with fitz.open(stream=content, filetype="pdf") as document:
        return [(index + 1, page.get_text("text")) for index, page in enumerate(document)]


def _extract_docx(content: bytes) -> list[tuple[int, str]]:
    document = Document(BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return [(1, "\n\n".join(parts))]


def extract_pages(filename: str, content: bytes) -> list[tuple[int, str]]:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Formato no soportado: {extension or 'sin extensión'}")
    if not content:
        raise ValueError("El archivo está vacío")
    return _extract_pdf(content) if extension == ".pdf" else _extract_docx(content)


def validate_uploads(files: list[dict[str, Any]]) -> None:
    if not files:
        raise ValueError("Seleccione al menos un documento")
    if len(files) > MAX_DOCUMENTS:
        raise ValueError(f"Se permiten máximo {MAX_DOCUMENTS} documentos por ejecución")
    total = 0
    for item in files:
        name = str(item.get("name") or "")
        content = bytes(item.get("content") or b"")
        extension = Path(name).suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise ValueError(f"Formato no soportado para {name}")
        if not content:
            raise ValueError(f"El archivo {name} está vacío")
        if len(content) > MAX_FILE_BYTES:
            raise ValueError(f"El archivo {name} supera 20 MB")
        total += len(content)
    if total > MAX_TOTAL_BYTES:
        raise ValueError("La carga total supera 50 MB")


def _safe_project_id(project_id: str) -> str:
    ascii_value = unicodedata.normalize("NFKD", str(project_id)).encode("ascii", "ignore").decode("ascii")
    normalized = re.sub(r"[^A-Za-z0-9_-]+", "-", ascii_value.strip()).strip("-")
    return normalized[:60] or "PROJECT-SESSION"


def ingest_project_documents(
    project_id: str,
    files: list[dict[str, Any]],
    encoder: Any,
) -> dict[str, Any]:
    """Extrae, limpia, fragmenta y vectoriza documentos sin persistirlos."""

    validate_uploads(files)
    safe_project = _safe_project_id(project_id)
    documents: list[dict[str, Any]] = []
    chunks: list[dict[str, Any]] = []
    errors: list[dict[str, str]] = []
    seen_hashes: set[str] = set()

    for item in files:
        filename = Path(str(item["name"])).name
        content = bytes(item["content"])
        content_hash = sha256(content).hexdigest()
        if content_hash in seen_hashes:
            errors.append({"filename": filename, "error": "Documento duplicado en la misma carga"})
            continue
        seen_hashes.add(content_hash)
        doc_id = f"{safe_project}-{content_hash[:12]}"
        try:
            pages = extract_pages(filename, content)
            document_chunks = 0
            extracted_characters = 0
            counter = 1
            for page_number, raw_text in pages:
                page_text = clean_text(raw_text)
                extracted_characters += len(page_text)
                for chunk_text in recursive_character_chunks(page_text):
                    chunks.append({
                        "project_id": safe_project,
                        "doc_id": doc_id,
                        "filename": filename,
                        "page": page_number,
                        "chunk_id": f"{doc_id}-RCHUNK-{counter:04d}",
                        "chunk_text": chunk_text,
                        "chunk_size_words": len(chunk_text.split()),
                        "chunk_size_chars": len(chunk_text),
                        "chunking_strategy": "recursive_character_splitter",
                        "chunk_size_param_chars": CHUNK_SIZE_CHARS,
                        "overlap_param_chars": CHUNK_OVERLAP_CHARS,
                    })
                    counter += 1
                    document_chunks += 1
            if not document_chunks:
                raise ValueError("No se encontró texto extraíble; puede ser un documento escaneado")
            documents.append({
                "project_id": safe_project,
                "doc_id": doc_id,
                "filename": filename,
                "extension": Path(filename).suffix.lower(),
                "size_bytes": len(content),
                "pages": len(pages),
                "extracted_characters": extracted_characters,
                "chunks": document_chunks,
                "status": "PROCESSED",
            })
        except Exception as exc:
            errors.append({"filename": filename, "error": str(exc)[:300]})

    if not chunks:
        raise ValueError("Ningún documento produjo texto utilizable")

    texts = [item["chunk_text"] for item in chunks]
    vectors = np.asarray(
        encoder.encode(
            texts,
            batch_size=16,
            show_progress_bar=False,
            normalize_embeddings=True,
            convert_to_numpy=True,
        ),
        dtype=np.float32,
    )
    if len(vectors) != len(chunks):
        raise ValueError("El número de embeddings no coincide con los chunks")
    for item, embedding in zip(chunks, vectors, strict=True):
        item["embedding"] = embedding

    return {
        "project_id": safe_project,
        "documents": documents,
        "chunks": chunks,
        "errors": errors,
        "summary": {
            "documents_received": len(files),
            "documents_processed": len(documents),
            "documents_rejected": len(errors),
            "pages": sum(item["pages"] for item in documents),
            "chunks": len(chunks),
            "embedding_dimensions": int(vectors.shape[1]) if vectors.ndim == 2 else 0,
            "status": "COMPLETED" if not errors else "COMPLETED_WITH_WARNINGS",
        },
        "configuration": {
            "chunking": "recursive_character_splitter",
            "chunk_size_chars": CHUNK_SIZE_CHARS,
            "chunk_overlap_chars": CHUNK_OVERLAP_CHARS,
            "embedding_model": EMBEDDING_MODEL,
            "normalize_embeddings": True,
            "persistence": "session_memory_only",
        },
    }
